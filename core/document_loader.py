"""Local extraction for PDF and DOCX workshop documents."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document

from core.text_cleaner import clean_text


class DocumentLoadError(ValueError):
    """Raised when a supported document cannot be read meaningfully."""


@dataclass(frozen=True)
class DocumentPage:
    text: str
    page_number: int | None = None


@dataclass(frozen=True)
class LoadedDocument:
    filename: str
    pages: list[DocumentPage]
    character_count: int


def load_document(filename: str, content: bytes) -> LoadedDocument:
    """Extract text from a PDF or DOCX byte stream while preserving source pages."""
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                pages = [
                    DocumentPage(clean_text(page.get_text("text")), index)
                    for index, page in enumerate(pdf, start=1)
                ]
        elif suffix == ".docx":
            document = Document(BytesIO(content))
            paragraphs = [clean_text(paragraph.text) for paragraph in document.paragraphs]
            text = "\n\n".join(paragraph for paragraph in paragraphs if paragraph)
            for table in document.tables:
                for row in table.rows:
                    cells = [clean_text(cell.text) for cell in row.cells]
                    row_text = " | ".join(cell for cell in cells if cell)
                    if row_text:
                        text = f"{text}\n{row_text}" if text else row_text
            pages = [DocumentPage(text, None)] if text else []
        else:
            raise DocumentLoadError("Please upload a PDF or DOCX file.")
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"Could not read {filename}: {exc}") from exc

    pages = [page for page in pages if page.text]
    character_count = sum(len(page.text) for page in pages)
    if not pages or character_count == 0:
        raise DocumentLoadError(
            f"No selectable text was found in {filename}. It may be image-only; "
            "run OCR first or choose a text-based document."
        )
    return LoadedDocument(filename, pages, character_count)
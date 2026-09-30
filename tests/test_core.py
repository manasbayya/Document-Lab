from io import BytesIO

import numpy as np
import pytest

from core.chunker import DocumentChunk, create_chunks
from core.document_loader import DocumentLoadError, DocumentPage, LoadedDocument, load_document
from core.embeddings import EmbeddingService
from core.model_selector import select_model
from core.retriever import retrieve
from core.text_cleaner import clean_text


def test_text_cleaning_preserves_structure_and_punctuation():
    assert clean_text("  Keep punctuation!\r\n\r\n\r\n  and words.\x00 ") == "Keep punctuation!\n\nand words."


def test_chunking_preserves_source_page_and_overlap():
    document = LoadedDocument("notes.pdf", [DocumentPage("word " * 100, 3)], 500)
    chunks = create_chunks(document, chunk_size=160, overlap=30)
    assert len(chunks) > 1
    assert all(chunk.source == "notes.pdf" and chunk.page_number == 3 for chunk in chunks)
    assert len({chunk.chunk_id for chunk in chunks}) == len(chunks)
    assert chunks[0].text.split()[-5:] == chunks[1].text.split()[:5]


def test_chunk_configuration_is_validated():
    document = LoadedDocument("notes.docx", [DocumentPage("text", None)], 4)
    with pytest.raises(ValueError):
        create_chunks(document, chunk_size=100, overlap=100)


def test_embeddings_are_normalized_and_empty_input_is_supported():
    class FakeModel:
        def encode(self, texts, **kwargs):
            assert kwargs["normalize_embeddings"] is True
            return np.array([[3.0, 4.0] for _ in texts]) / 5.0

    service = EmbeddingService(model=FakeModel())
    vectors = service.encode(["a", "b"])
    assert np.allclose(np.linalg.norm(vectors, axis=1), 1.0)
    assert service.encode([]).shape == (0, 0)


def test_embedding_model_loading_is_local_only(monkeypatch):
    import sys
    from types import SimpleNamespace

    calls = {}

    def fake_sentence_transformer(model_name, **kwargs):
        calls.update(model_name=model_name, **kwargs)
        return object()

    monkeypatch.setitem(
        sys.modules,
        "sentence_transformers",
        SimpleNamespace(SentenceTransformer=fake_sentence_transformer),
    )
    EmbeddingService()._load_model()
    assert calls["local_files_only"] is True
    assert calls["device"] == "cpu"


def test_cosine_retrieval_ranks_filters_and_limits_results():
    chunks = [
        DocumentChunk("a", "eligibility admission requirement", "guide.pdf", 1),
        DocumentChunk("b", "eligibility admission requirement", "guide.pdf", 1),
        DocumentChunk("c", "weather forecast rain", "guide.pdf", 2),
    ]
    vectors = np.array([[1.0, 0.0], [1.0, 0.0], [0.0, 1.0]])
    results = retrieve(np.array([1.0, 0.0]), vectors, chunks, top_k=3, threshold=0.8)
    assert [result.chunk.chunk_id for result in results] == ["b"]
    assert results[0].score == pytest.approx(1.0)


def test_model_selection_prefers_qwen_then_falls_back_to_llama():
    assert select_model(["llama3.2:3b", "qwen2.5:3b"]) == ("qwen2.5:3b", "Qwen")
    assert select_model(["llama3.1:8b"]) == ("llama3.1:8b", "Llama")
    assert select_model(["mistral:7b"]) == (None, None)


def test_empty_pdf_and_docx_documents_fail_gracefully():
    import pymupdf
    from docx import Document

    pdf = pymupdf.open()
    pdf.new_page()
    with pytest.raises(DocumentLoadError, match="No selectable text"):
        load_document("blank.pdf", pdf.tobytes())
    docx_stream = BytesIO()
    Document().save(docx_stream)
    with pytest.raises(DocumentLoadError, match="No selectable text"):
        load_document("blank.docx", docx_stream.getvalue())


def test_docx_extraction_includes_paragraphs_and_table_text():
    from docx import Document

    document = Document()
    document.add_paragraph("Eligibility requires 60 credits.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Program"
    table.cell(0, 1).text = "Engineering"
    stream = BytesIO()
    document.save(stream)
    loaded = load_document("syllabus.docx", stream.getvalue())
    assert loaded.character_count > 0
    assert "60 credits" in loaded.pages[0].text
    assert "Engineering" in loaded.pages[0].text


def test_pdf_extraction_preserves_page_numbers():
    import pymupdf

    pdf = pymupdf.open()
    first = pdf.new_page()
    first.insert_text((72, 72), "Eligibility is based on completed credits.")
    second = pdf.new_page()
    second.insert_text((72, 72), "Applications close in June.")
    loaded = load_document("syllabus.pdf", pdf.tobytes())
    assert [page.page_number for page in loaded.pages] == [1, 2]
    assert "completed credits" in loaded.pages[0].text


def test_rag_skips_generation_when_no_chunk_passes_threshold(monkeypatch):
    import core.rag_pipeline as pipeline

    class FakeEmbedder:
        def encode(self, texts):
            return np.array([[1.0, 0.0]])

    chunks = [DocumentChunk("c1", "irrelevant", "notes.pdf", 1)]

    def fail_if_called(*args, **kwargs):
        raise AssertionError("Ollama must not be called without retrieved evidence")

    monkeypatch.setattr(pipeline, "generate_answer", fail_if_called)
    result = pipeline.answer_question(
        "question", chunks, np.array([[0.0, 1.0]]), FakeEmbedder(), "qwen2.5:3b", threshold=0.8
    )
    assert result.answer == pipeline.NOT_FOUND
    assert result.sources == []


def test_rag_sends_only_retrieved_context_to_generation(monkeypatch):
    import core.rag_pipeline as pipeline

    class FakeEmbedder:
        def encode(self, texts):
            return np.array([[1.0, 0.0]])

    chunks = [
        DocumentChunk("relevant", "Eligibility requires 60 credits.", "guide.pdf", 2),
        DocumentChunk("other", "Rainfall is measured monthly.", "guide.pdf", 3),
    ]
    captured = {}

    def fake_generate(model, question, context, **kwargs):
        captured.update(model=model, question=question, context=context)
        return "At least 60 credits are required."

    monkeypatch.setattr(pipeline, "generate_answer", fake_generate)
    result = pipeline.answer_question(
        "Eligibility?", chunks, np.array([[1.0, 0.0], [0.0, 1.0]]), FakeEmbedder(),
        "qwen2.5:3b", top_k=1, threshold=0.5
    )
    assert result.answer == "At least 60 credits are required."
    assert "Eligibility requires 60 credits." in captured["context"]
    assert "Rainfall is measured monthly." not in captured["context"]
    assert captured["model"] == "qwen2.5:3b"